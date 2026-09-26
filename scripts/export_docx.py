#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PAPER_FINAL_EXPORT — Markdown -> DOCX (editable Word).
Only transforms layout; does NOT alter any number, claim, formula, or model content.
Figures: SVG rendered to PNG (2x) then embedded.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Mm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "paper" / "paper_v2_submission_candidate.md"
PNG_DIR = ROOT / "build" / "fig_png"


def set_cjk_font(run, size=None, bold=False):
    run.font.name = "Times New Roman"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), '宋体')
    if size:
        run.font.size = Pt(size)
    if bold:
        run.font.bold = True


def add_rich_text(paragraph, text, base_size=10.5, base_bold=False):
    """Parse **bold** / *italic* / `code` inline."""
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2])
            set_cjk_font(r, base_size, bold=True)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            set_cjk_font(r, base_size)
            r.font.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(max(base_size - 1.5, 7))
        else:
            r = paragraph.add_run(tok)
            set_cjk_font(r, base_size, bold=base_bold)


def md_to_docx(md: str, doc: Document):
    lines = md.split("\n")
    i, n = 0, len(lines)

    while i < n:
        line = lines[i]

        # fenced code block -> formula (centered, consolas)
        if line.strip().startswith("```"):
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            code = "\n".join(buf).strip("\n")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for j, cl in enumerate(code.split("\n")):
                r = p.add_run(cl)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                if j < len(code.split("\n")) - 1:
                    r.add_break()
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            lvl = len(m.group(1))
            text = m.group(2).strip()
            if lvl == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_rich_text(p, text, base_size=16, base_bold=True)
            else:
                h = doc.add_heading(level=min(lvl, 4))
                add_rich_text(h, text, base_size={2: 14, 3: 12, 4: 11}.get(lvl, 11), base_bold=True)
            i += 1
            continue

        # image
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if m:
            src = m.group(2)
            fname = Path(src).stem  # e.g. F1_model_flow
            png = PNG_DIR / f"{fname}.png"
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if png.exists():
                p.add_run().add_picture(str(png), width=Inches(6.2))
            else:
                add_rich_text(p, f"[图缺失: {fname}]", base_bold=True)
            i += 1
            continue

        # table
        if line.strip().startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            body = [r for r in rows if not re.match(r"^\|[\s:\-|]+\|$", r)]
            if body:
                parsed = [[c.strip() for c in r.strip("|").split("|")] for r in body]
                ncol = len(parsed[0])
                tbl = doc.add_table(rows=len(parsed), cols=ncol)
                tbl.style = 'Table Grid'
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(parsed):
                    for ci in range(ncol):
                        cell_text = row[ci] if ci < len(row) else ""
                        cell = tbl.rows[ri].cells[ci]
                        cp = cell.paragraphs[0]
                        add_rich_text(cp, cell_text, base_size=8.5, base_bold=(ri == 0))
                # spacer after table
                doc.add_paragraph()
            continue

        # blockquote
        if line.strip().startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(6)
            for j, b in enumerate(buf):
                r = p.add_run(b)
                set_cjk_font(r, 10)
                r.font.italic = True
                if j < len(buf) - 1:
                    r.add_break()
            continue

        # horizontal rule
        if line.strip() == "---":
            i += 1
            continue

        # list
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            ordered = bool(re.match(r"\d+\.", m.group(2)))
            items = []
            while i < n:
                m2 = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if not m2:
                    break
                items.append(m2.group(3).strip())
                i += 1
            for it in items:
                style = 'List Number' if ordered else 'List Bullet'
                p = doc.add_paragraph(style=style)
                add_rich_text(p, it, base_size=10)
            continue

        # blank
        if line.strip() == "":
            i += 1
            continue

        # normal paragraph (accumulate consecutive text)
        para_lines = []
        while i < n and lines[i].strip() and not re.match(r"^(#{1,6}\s|\||```|>|!\[|\s*[-*]\s|\s*\d+\.\s|---$)", lines[i]):
            para_lines.append(lines[i].strip())
            i += 1
        if para_lines:
            p = doc.add_paragraph()
            add_rich_text(p, " ".join(para_lines), base_size=10.5)
        else:
            i += 1
    return doc


def main():
    md = SRC.read_text(encoding="utf-8")
    doc = Document()
    # page setup A4, margins ~2.2cm
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(22)
    sec.bottom_margin = Mm(22)
    sec.left_margin = Mm(22)
    sec.right_margin = Mm(22)

    md_to_docx(md, doc)

    out = ROOT / "build" / "数学建模论文_可编辑版_2026-09-27.docx"
    doc.save(str(out))
    print(f"DOCX written: {out}")


if __name__ == "__main__":
    main()
