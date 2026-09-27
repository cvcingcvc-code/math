# -*- coding: utf-8 -*-
"""
PPT_FINAL_HUMAN_POLISH_IMPLEMENTATION
把 10 页 PPTX 重建为「中文优先」路演版：
  - 中文 = 主语言 / 大字体
  - 英文 = 专业术语 / 小字体 / 括号辅助
  - 状态码 = 小字 / 页脚 / 二级标签（保留不删除）
只重写 ppt/slides/slideN.xml 与其 rels；母版 / 主题 / 媒体 / docProps 原样保留。
"""
import os, re, shutil, zipfile, hashlib, datetime

# 仓库根 = 本脚本上两级（presentation/final_defense_ppt/ -> presentation/ -> 仓库根）
REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

# 母版/主题/媒体基座：读自原始 roadshow PPTX（它保留母版、主题、内嵌 SVG 与 docProps）
SRC = os.path.join(REPO, "outputs", "education_ai_evidence_roadshow_final.pptx")
# 三个正式交付副本（内容一致）
PEER = os.path.join(REPO, "presentation", "final_defense_ppt", "教育AI证据可靠性_现场答辩终稿.pptx")
THIRD = os.path.join(REPO, "outputs", "submission_package_development_only", "slides", "final_roadshow.pptx")
BUILD = os.path.join(REPO, "presentation", "final_defense_ppt", "_build_tmp")
os.makedirs(BUILD, exist_ok=True)
OUT = os.path.join(BUILD, "polished.pptx")

# ---------- 版式常量 ----------
EMU = 914400
def E(v): return int(round(v * EMU))

# 配色（沿用原盘面）
NAVY   = "0E3F8C"   # 主标题深蓝
BLUE   = "1E4FA8"   # 强调蓝
INK    = "1A2230"   # 正文近黑
GRAY   = "4A5568"   # 次级文字
MUTED  = "8B97A8"   # 弱化文字
RED    = "D9534F"   # 警示
WHITE  = "FFFFFF"
PANEL1 = "F7F9FC"   # 浅面板
PANEL2 = "F0F5FC"   # 浅蓝面板
PANELR = "FDECEA"   # 浅红面板
BORDER = "D6DCE5"
FONT   = "Microsoft YaHei"

# 页脚内容基线
def esc(t):
    return (t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

def run(text, sz, color, bold=False, lang="zh-CN"):
    b = ' b="1"' if bold else ' b="0"'
    return ('<a:r><a:rPr lang="%s" sz="%d"%s dirty="0"><a:solidFill><a:srgbClr val="%s"/>'
            '</a:solidFill><a:latin typeface="%s"/><a:ea typeface="%s"/></a:rPr>'
            '<a:t>%s</a:t></a:r>' % (lang, int(round(sz * 100)), b, color, FONT, FONT, esc(text)))

def tb(idx, x, y, w, h, lines, align="l", spacing=105000, anchor=None):
    """lines: list of (list_of_run_strings) —— 每项是一段"""
    ps = []
    for runs in lines:
        ppr = '<a:pPr algn="%s"><a:lnSpc><a:spcPct val="%d"/></a:lnSpc></a:pPr>' % (align, spacing)
        ps.append("<a:p>" + ppr + "".join(runs) + "</a:p>")
    bodypr = '<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0"'
    if anchor:
        bodypr += ' anchor="%s"' % anchor
    bodypr += "/>"
    return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name=""/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm rot="0" flipH="0" flipV="0"><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>'
            '<p:txBody>%s<a:lstStyle/>%s</p:txBody></p:sp>'
            % (idx, E(x), E(y), E(w), E(h), bodypr, "".join(ps)))

def rect(idx, x, y, w, h, fill=None, line=None, lw=1.0):
    f = '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % fill if fill else "<a:noFill/>"
    if line:
        ln = ('<a:ln w="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:ln>'
              % (int(round(lw * 12700)), line))
    else:
        ln = "<a:ln><a:noFill/></a:ln>"
    return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name=""/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm rot="0" flipH="0" flipV="0"><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>%s%s</p:spPr>'
            '<p:txBody><a:bodyPr/><a:p/></p:txBody></p:sp>'
            % (idx, E(x), E(y), E(w), E(h), f, ln))

def pic(idx, x, y, w, h, rid="rId1"):
    """SVG-only blip，与原始盘面完全同构"""
    return ('<p:pic><p:nvPicPr><p:cNvPr id="%d" name=""/><p:cNvPicPr><a:picLocks noChangeAspect="1"/>'
            '</p:cNvPicPr><p:nvPr/></p:nvPicPr><p:blipFill><a:blip><a:extLst>'
            '<a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}"><asvg:svgBlip r:embed="%s"/></a:ext>'
            '</a:extLst></a:blip><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            '<p:spPr><a:xfrm rot="0" flipH="0" flipV="0"><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:ln/></p:spPr></p:pic>'
            % (idx, rid, E(x), E(y), E(w), E(h)))

NS = ('<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="asvg">')

def slide(shapes):
    head = (NS + '<p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/>'
            '</p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
            '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>')
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + head
            + "".join(shapes) + '</p:spTree></p:cSld>'
            '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>')

# ---------- 公共构件 ----------
def bg(i=2):
    return rect(i, 0, 0, 13.3333, 7.5, fill=WHITE)

def header(idx, title_runs, sub=None, big=False):
    """页眉：标题 + 副题 + 规则线。返回 (shapes, next_idx)"""
    s = []
    i = idx
    if big:
        s.append(rect(i, 0.42, 0.42, 12.50, 1.04, fill=BLUE)); i += 1
        s.append(tb(i, 0.65, 0.68, 12.05, 0.56, [title_runs])); i += 1
    else:
        s.append(tb(i, 0.42, 0.46, 12.50, 0.62, [title_runs])); i += 1
        if sub:
            s.append(tb(i, 0.42, 1.04, 12.50, 0.26, [sub])); i += 1
        s.append(rect(i, 0.42, 1.32, 12.50, 0.035, fill=BLUE)); i += 1
    return s, i

def footer(idx, left=None, page=None):
    s = []
    i = idx
    if left:
        s.append(tb(i, 0.42, 7.02, 8.60, 0.22, [left])); i += 1
    if page:
        s.append(tb(i, 12.42, 7.02, 0.50, 0.22, [[run(page, 10, NAVY)]], align="r")); i += 1
    return s, i

def status_line(idx, y=6.62):
    """状态码降为小字页脚（保留，不删除）"""
    s = []
    s.append(tb(idx, 0.42, y, 12.50, 0.24, [[
        run("状态：", 9.5, MUTED), run("AI 临时标注", 9.5, MUTED),
        run(" AI_PROVISIONAL", 8, MUTED, lang="en-US"),
        run(" · 仅开发阶段", 9.5, MUTED),
        run(" DEVELOPMENT_ONLY", 8, MUTED, lang="en-US"),
        run(" · 尚未人工验证", 9.5, MUTED),
        run(" NOT_HUMAN_VALIDATED", 8, MUTED, lang="en-US"),
    ]]))
    return s

# ---------- 页面构建 ----------
def page01():
    s = [bg()]
    i = 2
    s.append(rect(i, 0.42, 0.42, 12.50, 0.62, fill=BLUE)); i += 1
    s.append(tb(i, 0.62, 0.58, 12.10, 0.32, [[run("教育 AI 增量价值评价｜现场答辩", 15, WHITE, True)]])); i += 1
    s.append(tb(i, 0.98, 1.52, 11.40, 0.92, [[run("观察到的结果，不等于可靠证据", 34, NAVY, True)]])); i += 1
    s.append(tb(i, 0.98, 2.58, 11.40, 0.58, [[run("这个分数，到底有多少证据值得相信？", 21, INK, True)]])); i += 1
    s.append(tb(i, 0.98, 3.30, 11.40, 0.42, [[
        run("当证据缺失、冲突或可信度不一，如何避免把观察值直接当成可信结论？", 13.5, GRAY)]])); i += 1
    s.append(rect(i, 0.98, 3.98, 9.70, 1.42, fill=PANEL2)); i += 1
    s.append(rect(i, 0.98, 3.98, 0.055, 1.42, fill=BLUE)); i += 1
    s.append(tb(i, 1.28, 4.20, 9.20, 0.44, [[run("Score Stability ≠ Evidence Support Stability", 17, BLUE, True, lang="en-US")]])); i += 1
    s.append(tb(i, 1.28, 4.74, 9.20, 0.42, [[run("分数稳定，不等于证据支撑稳定", 19, INK, True)]])); i += 1
    s.append(tb(i, 0.98, 5.56, 11.40, 0.56, [
        [run("一句话：我们在补的不是一个更高的分数，", 14, INK)],
        [run("而是「这个分数有多少证据值得相信」。", 14, INK)],
    ])); i += 1
    s.append(tb(i, 0.98, 6.20, 11.40, 0.36, [[
        run("起名困难队", 18, NAVY, True),
        run("　　", 14, NAVY),
        run("汇报人：张镇源", 15, INK, True),
    ]])); i += 1
    s += status_line(i, y=6.66); i += 1
    s += footer(i, [run("路演版 · 简体中文", 9, MUTED)], "01 / 10")[0]
    return slide(s)

def page02():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("两个问题必须分开看", 30, NAVY, True)],
                  [run("表现有多高，不等于证据有多可信", 13, GRAY)])
    s += h
    s.append(rect(i, 0.42, 1.52, 7.29, 5.00, fill=PANEL1, line=BORDER)); i += 1
    s.append(pic(i, 0.62, 2.01, 6.89, 4.02)); i += 1
    s.append(rect(i, 8.00, 1.52, 4.79, 5.00, fill=PANEL2)); i += 1
    s.append(tb(i, 8.23, 1.76, 4.33, 0.32, [[run("两个概念", 19, NAVY, True)]])); i += 1
    s.append(tb(i, 8.23, 2.28, 4.33, 0.72, [
        [run("表现有多高", 19, INK, True)],
        [run("（Raw Signal）", 10.5, MUTED, lang="en-US")],
    ])); i += 1
    s.append(tb(i, 8.23, 3.10, 4.33, 0.72, [
        [run("证据有多可信", 19, INK, True)],
        [run("（Evidence Reliability）", 10.5, MUTED, lang="en-US")],
    ])); i += 1
    s.append(rect(i, 8.23, 3.94, 4.33, 1.20, fill=PANELR)); i += 1
    s.append(rect(i, 8.23, 3.94, 0.05, 1.20, fill=RED)); i += 1
    s.append(tb(i, 8.48, 4.12, 3.90, 0.86, [
        [run("高表现，", 17, RED, True)],
        [run("不一定有高证据支持。", 17, RED, True)],
    ])); i += 1
    s.append(tb(i, 8.23, 5.32, 4.33, 0.60, [
        [run("归一化分数不变，", 11, GRAY)],
        [run("不代表支持范围不变。", 11, GRAY)],
    ])); i += 1
    s.append(tb(i, 0.42, 6.68, 6.20, 0.22, [[run("冻结图 F2 · 原始表现 vs 调整后表现", 9.5, MUTED)]])); i += 1
    s += footer(i, None, "02 / 10")[0]
    return slide(s)

def page03():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("模型链路：从观察到决策", 30, NAVY, True)],
                  [run("证据可靠性如何进入评价", 13, GRAY)])
    s += h
    s.append(rect(i, 0.42, 1.45, 7.00, 4.30, fill=PANEL1, line=BORDER)); i += 1
    s.append(pic(i, 0.62, 1.65, 6.60, 3.85)); i += 1
    s.append(rect(i, 7.65, 1.45, 5.14, 4.30, fill=PANEL2)); i += 1
    s.append(tb(i, 7.88, 1.62, 4.68, 0.30, [[run("支持权重由四个因子决定", 15, NAVY, True)]])); i += 1
    s.append(tb(i, 7.88, 2.02, 4.68, 2.60, [
        [run("可判读性", 15, INK, True)],
        [run("× 标注置信度", 15, INK, True)],
        [run("× AI 提示风险", 15, INK, True)],
        [run("× 上下文风险", 15, INK, True)],
    ])); i += 1
    s.append(tb(i, 7.88, 4.72, 4.68, 0.80, [
        [run("（后两项为风险折减）", 9.5, MUTED)],
        [run("Observability · Confidence", 8.5, MUTED, lang="en-US")],
        [run("· Prompt risk · Context risk", 8.5, MUTED, lang="en-US")],
    ])); i += 1
    s.append(rect(i, 0.42, 5.90, 12.50, 0.78, fill=PANEL2)); i += 1
    s.append(tb(i, 0.62, 6.06, 12.10, 0.46, [[
        run("w", 13, NAVY, lang="en-US"), run("ᵢ", 9, NAVY), run(" = I(observableᵢ) · cᵢ · (1 − λprompt·pᵢ) · (1 − λcontext·tᵢ)", 13, NAVY, lang="en-US"),
        run("        ", 12, NAVY), run("Adjustedᵢ = Rawᵢ × wᵢ", 13, NAVY, lang="en-US"),
        run("        ", 12, NAVY), run("Ceff = Σwᵢ / N", 13, NAVY, lang="en-US"),
    ]])); i += 1
    s.append(tb(i, 0.42, 6.84, 9.00, 0.24, [[
        run("Reliability 是支持权重，不是概率，也不是因果系数。", 10.5, GRAY)]])); i += 1
    s += footer(i, None, "03 / 10")[0]
    return slide(s)

def page04():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("70 条开发记录的证据分布", 30, NAVY, True)],
                  [run("可判读与不可判读必须分别报告", 13, GRAY)])
    s += h
    cards = [
        ("70", "70 条开发记录", "AI provisional records", NAVY),
        ("16", "16 条可判读证据", "readable evidence", NAVY),
        ("54", "54 条无有效证据", "NO_EFFECTIVE_EVIDENCE", RED),
    ]
    x = 0.42
    for num, zh, en, col in cards:
        s.append(rect(i, x, 1.55, 4.00, 2.55, fill=PANEL1, line=BORDER)); i += 1
        s.append(tb(i, x + 0.30, 1.80, 3.40, 1.00, [[run(num, 54, col, True)]])); i += 1
        s.append(tb(i, x + 0.30, 2.95, 3.40, 0.38, [[run(zh, 18, INK, True)]])); i += 1
        s.append(tb(i, x + 0.30, 3.42, 3.40, 0.34, [[run("（%s）" % en, 10, MUTED, lang="en-US")]])); i += 1
        x += 4.25
    s.append(rect(i, 0.42, 4.32, 12.50, 1.62, fill=PANELR)); i += 1
    s.append(rect(i, 0.42, 4.32, 0.055, 1.62, fill=RED)); i += 1
    s.append(tb(i, 0.75, 4.52, 11.90, 0.60, [[run("54 不是 54 个零分。", 26, RED, True)]])); i += 1
    s.append(tb(i, 0.75, 5.22, 11.90, 0.52, [[
        run("而是 54 条没有足够证据进行可靠数值判断的记录。", 18, INK, True)]])); i += 1
    s.append(tb(i, 0.42, 6.14, 12.50, 0.50, [
        [run("其中 19 条无可判读证据、35 条证据不确定；二者合计 54，均不填 0。", 11, GRAY)],
    ])); i += 1
    s += footer(i, [run("数据状态：AI 临时标注 · 仅开发阶段", 9.5, MUTED)], "04 / 10")[0]
    return slide(s)

def page05():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("分数没有变化，但支撑这个分数的证据明显减少", 26, NAVY, True)],
                  [run("这是全场最重要的一页：分数稳定 ≠ 证据支撑稳定", 13, RED, True)])
    s += h
    s.append(rect(i, 0.42, 1.52, 7.00, 4.22, fill=PANEL1, line=BORDER)); i += 1
    s.append(pic(i, 0.62, 1.72, 6.60, 3.85)); i += 1
    s.append(rect(i, 7.65, 1.52, 5.14, 4.22, fill=PANEL2)); i += 1
    s.append(tb(i, 7.88, 1.70, 4.68, 0.30, [[run("三项按重要性排序", 14, NAVY, True)]])); i += 1
    s.append(tb(i, 7.88, 2.10, 4.68, 0.28, [[run("① 有效支持", 13, GRAY)]])); i += 1
    s.append(tb(i, 7.88, 2.38, 4.68, 0.42, [[run("16 → 12 → 6 → 6", 19, BLUE, True, lang="en-US")]])); i += 1
    s.append(tb(i, 7.88, 2.92, 4.68, 0.28, [[run("② 有效覆盖率", 13, GRAY)]])); i += 1
    s.append(tb(i, 7.88, 3.20, 4.68, 0.42, [[run("22.86% → 8.57%", 19, BLUE, True, lang="en-US")]])); i += 1
    s.append(tb(i, 7.88, 3.64, 4.68, 0.40, [
        [run("0.228571 → 0.171429", 8.5, MUTED, lang="en-US")],
        [run("→ 0.085714 → 0.085714", 8.5, MUTED, lang="en-US")],
    ])); i += 1
    s.append(tb(i, 7.88, 4.06, 4.68, 0.28, [[run("③ 分数", 13, GRAY)]])); i += 1
    s.append(tb(i, 7.88, 4.34, 4.68, 0.42, [[run("3.5625 → 3.5625", 19, GRAY, True, lang="en-US")]])); i += 1
    s.append(tb(i, 7.88, 4.92, 4.68, 0.66, [
        [run("支持与覆盖率下降，", 10.5, GRAY)],
        [run("分数保持不变。", 10.5, GRAY)],
    ])); i += 1
    s.append(rect(i, 0.42, 5.90, 12.50, 0.56, fill=PANELR)); i += 1
    s.append(rect(i, 0.42, 5.90, 0.055, 0.56, fill=RED)); i += 1
    s.append(tb(i, 0.75, 6.04, 11.90, 0.32, [[run("分数稳定 ≠ 证据支撑稳定", 16, RED, True)]])); i += 1
    s.append(tb(i, 0.42, 6.58, 12.50, 0.26, [[
        run("ABL / HOT / Gap = 3.5625 / 0.375 / 1.125", 10.5, GRAY, lang="en-US"),
        run("（分数不变，不是提升）", 10.5, GRAY)]])); i += 1
    s.append(tb(i, 0.42, 6.84, 6.00, 0.22, [[run("冻结图 F5 · 组件消融与拒判边界", 9.5, MUTED)]])); i += 1
    s += footer(i, None, "05 / 10")[0]
    return slide(s)

def page06():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("五种方法，验证同一个问题", 30, NAVY, True)],
                  [run("当分数看起来稳定时，证据支持是否也稳定？", 14, INK, True)])
    s += h
    cards = [
        ("01", "基线对比", "Baseline", "不建模可靠性会怎样"),
        ("02", "组件消融", "Ablation", "支持量是否真被组件改变"),
        ("03", "参数敏感性", "Sensitivity", "结论是否只是碰巧"),
        ("04", "失败案例", "Failure Case", "高 Raw 是否仍被接受"),
        ("05", "外部结构迁移", "External Transfer", "结构能否跨域运行"),
    ]
    x = 0.42
    for n, zh, en, desc in cards:
        s.append(rect(i, x, 1.50, 2.33, 1.90, fill=PANEL1, line=BORDER)); i += 1
        s.append(tb(i, x + 0.18, 1.62, 2.00, 0.24, [[run(n, 10, MUTED, lang="en-US")]])); i += 1
        s.append(tb(i, x + 0.18, 1.90, 2.00, 0.34, [[run(zh, 15, BLUE, True)]])); i += 1
        s.append(tb(i, x + 0.18, 2.28, 2.00, 0.24, [[run("（%s）" % en, 9, MUTED, lang="en-US")]])); i += 1
        s.append(tb(i, x + 0.18, 2.66, 2.02, 0.60, [[run(desc, 11, INK)]])); i += 1
        x += 2.50
    # 参数敏感性证据卡（补回锁定数字）
    s.append(rect(i, 0.42, 3.46, 7.90, 3.10, fill=PANEL2)); i += 1
    s.append(rect(i, 0.42, 3.46, 0.055, 3.10, fill=BLUE)); i += 1
    s.append(tb(i, 0.70, 3.62, 7.40, 0.28, [[run("参数敏感性证据卡（Sensitivity）", 13, NAVY, True)]])); i += 1
    s.append(tb(i, 0.70, 3.98, 7.40, 0.44, [[
        run("693 = 21 × 11 × 3", 21, BLUE, True, lang="en-US"),
        run("     ", 14, NAVY),
        run("660 defined", 14, INK, lang="en-US"),
        run("     ", 12, NAVY),
        run("33 undefined", 14, INK, lang="en-US"),
    ]])); i += 1
    s.append(tb(i, 0.70, 4.52, 7.40, 0.34, [[
        run("扰动幅度：±10% / ±20%", 13, INK, lang="zh-CN"),
    ]])); i += 1
    s.append(tb(i, 0.70, 4.94, 7.40, 0.40, [[
        run("0 decision-state flips", 19, RED, True, lang="en-US"),
        run("（决策状态零翻转）", 12, RED),
    ]])); i += 1
    s.append(tb(i, 0.70, 5.50, 7.40, 0.90, [
        [run("五条路径不是五个新主张，", 14, INK, True)],
        [run("而是同一个可审查问题的五种压力测试。", 14, INK, True)],
        [run("每条路径统一报告：为什么问 → 怎么设 → 结果 → 支持什么 → 不支持什么", 10.5, GRAY)],
    ])); i += 1
    s.append(pic(i, 8.60, 3.62, 4.35, 2.80)); i += 1
    s += footer(i, [run("不新增实验解释，只收束既有冻结证据", 9.5, MUTED)], "06 / 10")[0]
    return slide(s)

def page07():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("高 Raw，不代表高可信", 32, NAVY, True)],
                  [run("最困难的记录，正好暴露支持与拒判的边界", 13, GRAY)])
    s += h
    s.append(rect(i, 0.42, 1.52, 7.00, 5.06, fill=PANEL1, line=BORDER)); i += 1
    s.append(pic(i, 0.62, 1.72, 6.60, 3.85)); i += 1
    s.append(tb(i, 0.62, 5.72, 6.60, 0.70, [
        [run("冻结图 F3 · 证据退化与拒判边界", 9.5, MUTED)],
        [run("受控改变单个证据字段时的机械响应，不是真实观测。", 9.5, MUTED)],
    ])); i += 1
    s.append(rect(i, 7.65, 1.52, 5.14, 2.32, fill=PANELR)); i += 1
    s.append(rect(i, 7.65, 1.52, 0.055, 2.32, fill=RED)); i += 1
    s.append(tb(i, 7.92, 1.70, 4.64, 0.34, [[run("P108", 18, INK, True, lang="en-US")]])); i += 1
    s.append(tb(i, 7.92, 2.10, 4.64, 1.24, [
        [run("原始表现 = L6", 15, INK, True), run("（最高层）", 10, MUTED)],
        [run("证据可靠性 = 0.375", 15, INK, True)],
        [run("调整后贡献 = 2.25", 15, INK, True)],
    ])); i += 1
    s.append(tb(i, 7.92, 3.42, 4.64, 0.34, [[
        run("→ 支持不足，暂缓判断（ABSTAIN）", 13, RED, True)]])); i += 1
    s.append(rect(i, 7.65, 4.00, 5.14, 1.30, fill=PANEL2)); i += 1
    s.append(tb(i, 7.92, 4.16, 4.64, 1.02, [
        [run("P072 / P035", 16, INK, True, lang="en-US")],
        [run("→ 无有效证据", 13.5, INK)],
        [run("→ 不填 0", 13.5, RED, True)],
    ])); i += 1
    s.append(tb(i, 7.65, 5.46, 5.14, 0.56, [
        [run("没有证据，不等于表现为零。", 17, INK, True)],
    ])); i += 1
    s.append(tb(i, 7.65, 6.06, 5.14, 0.50, [
        [run("同类低支持案例：P105（L2 · R=0.375 · 贡献 0.75）", 9.5, MUTED, lang="zh-CN")],
        [run("LOW_SUPPORT", 8.5, MUTED, lang="en-US")],
    ])); i += 1
    s += footer(i, None, "07 / 10")[0]
    return slide(s)

def page08():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("补充压力测试与结构迁移", 30, NAVY, True)],
                  [run("支持边界与支持范围同样重要", 13, GRAY)])
    s += h
    s.append(rect(i, 0.42, 1.46, 12.50, 0.60, fill=PANEL2)); i += 1
    s.append(rect(i, 0.42, 1.46, 0.055, 0.60, fill=BLUE)); i += 1
    s.append(tb(i, 0.75, 1.60, 11.90, 0.36, [[
        run("结构可以迁移，预测优势没有被证明。", 19, INK, True)]])); i += 1
    s.append(rect(i, 0.42, 2.20, 6.05, 2.16, fill=PANEL1, line=BORDER)); i += 1
    s.append(tb(i, 0.65, 2.30, 5.60, 0.26, [[
        run("机器学习对照（ML Challenger）", 13, NAVY, True)]])); i += 1
    s.append(tb(i, 0.65, 2.58, 5.60, 0.26, [[
        run("核心结论：没有足够证据选出 winner。", 12.5, INK, True)]])); i += 1
    s.append(tb(i, 0.65, 2.88, 5.60, 1.34, [
        [run("表面判别指标 0.876 / 0.919 / 0.922", 10, GRAY, lang="zh-CN")],
        [run("RF 跨学期 ≈ 0.588；缺失后约 0.71–0.81", 10, GRAY, lang="zh-CN")],
        [run("但 positive 仅 16 条且为 AI 代理标签", 10, GRAY)],
        [run("INSUFFICIENT_FOR_STRONG_ML_CLAIM", 8.5, RED, lang="en-US")],
    ])); i += 1
    s.append(rect(i, 6.68, 2.20, 6.24, 2.16, fill=PANEL2)); i += 1
    s.append(tb(i, 6.91, 2.34, 5.80, 0.30, [[
        run("外部结构迁移（External Structural Transfer）", 14, NAVY, True)]])); i += 1
    s.append(tb(i, 6.91, 2.70, 5.80, 0.30, [[
        run("BTC 历史数据，仅做结构迁移验证（n = 1698）。", 12, INK)]])); i += 1
    s.append(tb(i, 6.91, 3.06, 5.80, 1.16, [
        [run("Gate 1：支持", 13, INK, True)],
        [run("Gate 2：支持", 13, INK, True)],
        [run("Gate 3：当前证据不支持", 13, RED, True)],
        [run("低可靠误差 0.5248 vs 高可靠误差 0.5277", 9, MUTED)],
    ])); i += 1
    s.append(pic(i, 4.44, 4.52, 4.46, 2.60)); i += 1
    s += footer(i, [run("负结果必须保留：Gate 3 不支持", 9.5, MUTED)], "08 / 10")[0]
    return slide(s)

def page09():
    s = [bg()]
    i = 2
    h, i = header(i,
                  [run("当前证据边界", 30, NAVY, True)],
                  [run("现在能说什么，不能说什么", 13, GRAY)])
    s += h
    s.append(rect(i, 0.42, 1.50, 6.05, 3.50, fill=PANEL2)); i += 1
    s.append(tb(i, 0.68, 1.68, 5.55, 0.34, [[run("【现在能说】", 16, NAVY, True)]])); i += 1
    s.append(tb(i, 0.68, 2.16, 5.55, 2.60, [
        [run("— 结构性结论", 14, INK)],
        [run("— 描述性结论", 14, INK)],
        [run("— 证据支持变化", 14, INK)],
        [run("— 拒判机制行为", 14, INK)],
    ])); i += 1
    s.append(rect(i, 6.88, 1.50, 6.04, 3.50, fill=PANELR)); i += 1
    s.append(tb(i, 7.14, 1.68, 5.54, 0.34, [[run("【现在不能说】", 16, RED, True)]])); i += 1
    s.append(tb(i, 7.14, 2.16, 5.54, 2.80, [
        [run("— AI 提升学习效果", 14, INK)],
        [run("— 因果 AI 增量", 14, INK)],
        [run("— 准确率提升", 14, INK)],
        [run("— 交易优势", 14, INK)],
        [run("— 通用泛化", 14, INK)],
    ])); i += 1
    s.append(rect(i, 0.42, 5.20, 12.50, 1.06, fill=PANEL1, line=BORDER)); i += 1
    s.append(tb(i, 0.65, 5.34, 12.05, 0.30, [[
        run("AI 临时标注，尚未人工验证", 12, GRAY)]])); i += 1
    s.append(tb(i, 0.65, 5.72, 12.05, 0.34, [[
        run("R1 = 0/70  ·  R2 = 0/70", 11, INK, lang="en-US"),
        run("   ·   ", 10, MUTED),
        run("Formal Gate = NOT_RUN", 11, INK, lang="en-US"),
        run("   ·   ", 10, MUTED),
        run("AI Increment = NOT_SUPPORTED", 11, RED, lang="en-US"),
    ]])); i += 1
    s += footer(i, [run("边界不是附注，而是结果可信度的一部分", 9.5, MUTED)], "09 / 10")[0]
    return slide(s)

def page10():
    s = [bg()]
    i = 2
    h, i = header(i, [run("结论与现场 Demo", 25, WHITE, True)], big=True)
    s += h
    cards = [
        ("01", "问题", "观测结果能否直接当作可靠证据？"),
        ("02", "模型", "原始表现 → 证据可靠性 → 证据支持 → 决策"),
        ("03", "结果", "分数不变；支持 16 → 6"),
    ]
    x = 0.42
    for n, zh, desc in cards:
        s.append(rect(i, x, 1.72, 3.90, 1.86, fill=PANEL1, line=BORDER)); i += 1
        s.append(tb(i, x + 0.24, 1.88, 3.42, 0.26, [[run(n, 10, MUTED, lang="en-US")]])); i += 1
        s.append(tb(i, x + 0.24, 2.16, 3.42, 0.34, [[run(zh, 17, BLUE, True)]])); i += 1
        s.append(tb(i, x + 0.24, 2.62, 3.42, 0.82, [[run(desc, 12, INK)]])); i += 1
        x += 4.30
    s.append(rect(i, 0.42, 3.82, 12.50, 1.52, fill=PANEL2)); i += 1
    s.append(rect(i, 0.42, 3.82, 0.055, 1.52, fill=BLUE)); i += 1
    s.append(tb(i, 0.78, 4.04, 11.80, 1.10, [
        [run("我们真正增加的，不是一个更高的分数，", 21, INK, True)],
        [run("而是让模型知道：什么时候证据不够。", 21, NAVY, True)],
    ])); i += 1
    s.append(tb(i, 0.42, 5.52, 12.50, 0.60, [
        [run("现场 Demo：离线展示页 + 冻结图 F1–F6", 14, BLUE, True)],
        [run("（不依赖网络，断网可讲）", 10, MUTED)],
    ])); i += 1
    s += status_line(i, y=6.28); i += 1
    s += footer(i, [run("起名困难队 · 张镇源", 9, MUTED)], "10 / 10")[0]
    return slide(s)

PAGES = {1: page01, 2: page02, 3: page03, 4: page04, 5: page05,
         6: page06, 7: page07, 8: page08, 9: page09, 10: page10}

# 幻灯片 -> 内嵌图（media/imageN.svg）
IMG = {2: "image1.svg", 3: "image2.svg", 5: "image5.svg",
       6: "image4.svg", 7: "image3.svg", 8: "image6.svg"}

REL_LAYOUT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
              '<Relationship Id="rId0" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml" />'
              '</Relationships>')

def rels_for(n):
    base = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId0" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml" />')
    if n in IMG:
        base += ('<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/%s" />' % IMG[n])
    return base + "</Relationships>"

# ---------- 打包 ----------
src_zip = zipfile.ZipFile(SRC, "r")
names = src_zip.namelist()
new_slides = {("ppt/slides/slide%d.xml" % n): PAGES[n]() for n in range(1, 11)}
new_rels = {("ppt/slides/_rels/slide%d.xml.rels" % n): rels_for(n) for n in range(1, 11)}

if os.path.exists(OUT):
    os.remove(OUT)
with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zo:
    for item in src_zip.infolist():
        data = src_zip.read(item.filename)
        if item.filename in new_slides:
            data = new_slides[item.filename].encode("utf-8")
        elif item.filename in new_rels:
            data = new_rels[item.filename].encode("utf-8")
        zo.writestr(item, data)
src_zip.close()

# 回读校验
from pptx import Presentation
chk = Presentation(OUT)
print("BUILD_OK slides=%d" % len(chk.slides))
print("OUT_SHA256 =", hashlib.sha256(open(OUT, "rb").read()).hexdigest())
for i, sl in enumerate(chk.slides, 1):
    chars = 0
    for sh in sl.shapes:
        if sh.has_text_frame:
            chars += len(sh.text_frame.text)
    print("  slide %02d  shapes=%2d  textchars=%d" % (i, len(sl.shapes), chars))

# 同步三个正式交付副本（SRC 基座本身也一并更新为中文优先版）
data = open(OUT, "rb").read()
for dst in (SRC, PEER, THIRD):
    open(dst, "wb").write(data)
    print("DEPLOYED %s  %s" % (hashlib.sha256(open(dst, "rb").read()).hexdigest()[:16], os.path.basename(dst)))
